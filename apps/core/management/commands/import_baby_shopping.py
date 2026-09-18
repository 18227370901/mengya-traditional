"""Django management command: 导入待产包 Excel 数据"""
import openpyxl
from django.core.management.base import BaseCommand
from apps.core.models import BabyShoppingItem


class Command(BaseCommand):
    help = "从 Excel 导入待产包参考数据（妈妈篇 + 宝宝篇）"

    def add_arguments(self, parser):
        parser.add_argument("file", type=str, help="Excel 文件路径")
        parser.add_argument("--clear", action="store_true", help="导入前清空旧数据")

    def handle(self, *args, **options):
        file_path = options["file"]
        clear = options.get("clear", False)

        if clear:
            BabyShoppingItem.objects.all().delete()
            self.stdout.write(self.style.WARNING("已清空旧数据"))

        wb = openpyxl.load_workbook(file_path, data_only=True)
        total = 0

        # Sheet → owner 映射
        sheet_owner = {"妈妈篇": "mom", "宝宝篇": "baby"}

        for sheet_name in wb.sheetnames:
            if sheet_name not in sheet_owner:
                continue
            owner = sheet_owner[sheet_name]
            ws = wb[sheet_name]
            current_category = None
            sort_order = 0

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True), 2):
                # 跳过全空行
                if all(v is None for v in row):
                    continue

                category = row[0] if len(row) > 0 else None
                name = row[1] if len(row) > 1 else None
                quantity = row[2] if len(row) > 2 else None
                unit = row[3] if len(row) > 3 else None
                remark = row[4] if len(row) > 4 else None
                # row[5] = 购买情况 (skip)
                unit_price = row[6] if len(row) > 6 else None
                total_price = row[7] if len(row) > 7 else None
                image_formula = row[8] if len(row) > 8 else None
                extra_image_formula = row[9] if len(row) > 9 else None

                # Update category if not None
                if category:
                    current_category = str(category).strip()

                # Skip if no item name
                if not name:
                    continue

                # Convert quantity to string
                if quantity is not None:
                    quantity = str(quantity).strip()
                else:
                    quantity = "1"

                # Convert unit
                if unit is not None:
                    unit = str(unit).strip()
                else:
                    unit = ""

                # Parse prices
                def parse_price(val):
                    if val is None or val == "":
                        return None
                    try:
                        return float(val)
                    except (ValueError, TypeError):
                        return None

                up = parse_price(unit_price)
                tp = parse_price(total_price)

                # Parse image URL from DISPIMG formula
                def parse_image(formula):
                    if not formula:
                        return None
                    s = str(formula)
                    if "DISPIMG" in s:
                        # Extract ID from =DISPIMG("ID_XXX",1)
                        import re
                        m = re.search(r'DISPIMG\("([^"]+)"', s)
                        if m:
                            return f"excel://image/{m.group(1)}"
                    return None

                img = parse_image(image_formula)
                img2 = parse_image(extra_image_formula)

                # Ensure remark is string
                if remark:
                    remark = str(remark).strip()
                else:
                    remark = ""

                BabyShoppingItem.objects.create(
                    owner=owner,
                    category=current_category or "其他",
                    name=str(name).strip(),
                    quantity=quantity,
                    unit=unit,
                    unit_price=up,
                    total_price=tp,
                    remark=remark,
                    image_url=img,
                    extra_image_url=img2,
                    sort_order=sort_order,
                )
                sort_order += 1
                total += 1

        self.stdout.write(
            self.style.SUCCESS(f"导入完成！共导入 {total} 条待产包参考数据")
        )
